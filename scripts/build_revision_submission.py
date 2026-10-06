"""Build editable major-revision documents without overwriting the submission.

Run with the task-owned document environment (python-docx). Pandoc is required.
Numerical inputs are read only. PDF rendering is a separate render_docx.py step.
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REV = ROOT / "revision/major_revision_20261006"
DEFAULT_OUT = ROOT / "submission/bmc_bioinformatics_major_revision_20261007"


def format_document(path: Path, main: bool) -> None:
    document = Document(path)
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(.72)
        section.left_margin = section.right_margin = Inches(.72)
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        footer.add_run("Page ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        if main:
            numbering = OxmlElement("w:lnNumType")
            numbering.set(qn("w:countBy"), "5")
            numbering.set(qn("w:restart"), "continuous")
            section._sectPr.append(numbering)
    for style in document.styles:
        if style.type == 1:
            style.font.name = "Times New Roman"
            style.font.size = Pt(11)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 2.0 if main else 1.15
            if style.name.startswith("Heading"):
                style.font.bold = True
                style.font.size = Pt(13 if style.name in ("Heading 1", "Heading 2") else 11)
                style.paragraph_format.keep_with_next = True
    for table in document.tables:
        table.autofit = True
        if main and len(table.columns) == 9:
            table.autofit = False
            widths = [1.35] + [(6.83 - 1.35) / 8] * 8
            for column, width in zip(table.columns, widths):
                column.width = Inches(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Inches(width)
        for i, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            if i == 0:
                repeat = OxmlElement("w:tblHeader")
                row._tr.get_or_add_trPr().append(repeat)
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(4)
                    for run in paragraph.runs:
                        run.font.size = Pt(8 if len(row.cells) > 5 else 9)
    paragraphs = document.paragraphs
    for i, paragraph in enumerate(paragraphs):
        if re.match(r"Table \d+\.", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
            if i + 1 < len(paragraphs):
                paragraphs[i + 1].paragraph_format.keep_with_next = True
        if paragraph._p.xpath('.//w:drawing'):
            paragraph.paragraph_format.keep_with_next = True
    document.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUT)
    parser.add_argument("--figures-dir", type=Path, required=True)
    parser.add_argument("--only", choices=("Manuscript", "Supplementary_Methods", "Response_to_Reviewers", "Cover_Letter"))
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    cache = ROOT / "data/cache/revision-writing-20261007/document-inputs"
    cache.mkdir(parents=True, exist_ok=True)
    sources = {
        "Manuscript": ROOT / "manuscript/draft.md",
        "Supplementary_Methods": ROOT / "manuscript/supplementary_methods.md",
        "Response_to_Reviewers": REV / "rebuttal.md",
        "Cover_Letter": REV / "cover_letter_revision.md",
    }
    for name, source in sources.items():
        if args.only and args.only != name:
            continue
        text = source.read_text()
        if name == "Manuscript":
            lines = []
            for line in text.splitlines():
                if re.fullmatch(r"\[Figure \d+ near here\]", line):
                    continue
                match = re.match(r"\*\*Figure (\d+)\.", line)
                if match:
                    figures = sorted(args.figures_dir.glob(f"fig{match[1]}_*.png"))
                    if len(figures) != 1:
                        raise ValueError(f"Expected one image for figure {match[1]}: {figures}")
                    lines.extend([f"![]({figures[0].resolve()}){{width=6.7in}}", ""])
                lines.append(line)
            text = "\n".join(lines)
        staged = cache / f"{name}.md"
        staged.write_text(text)
        output = args.output_dir / f"{name}.docx"
        subprocess.run(["pandoc", str(staged), "--syntax-highlighting=none", "-o", str(output)], check=True)
        format_document(output, name == "Manuscript")
    if args.only:
        print(f"Built {args.only}.docx")
        return
    original = ROOT / "submission/bmc_bioinformatics"
    for n in (1, 2, 4, 5, 7, 8, 9):
        files = list(original.glob(f"Additional_file_{n}_*"))
        if len(files) != 1:
            raise ValueError(files)
        shutil.copy2(files[0], args.output_dir / files[0].name)
    for source, name in [
        (sources["Supplementary_Methods"], "Additional_file_3_Supplementary_Methods.txt"),
        (ROOT / "data/report/exploration/expression_shortlist_report.md", "Additional_file_6_Expression_Shortlist_Report.txt"),
    ]:
        subprocess.run(["pandoc", str(source), "-t", "plain", "--wrap=none", "-o", str(args.output_dir / name)], check=True)
    print(f"Built four DOCX documents and nine additional files in {args.output_dir}")


if __name__ == "__main__":
    main()
