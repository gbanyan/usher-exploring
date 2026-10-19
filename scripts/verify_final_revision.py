"""Verify the current author-review package against its sources and invariants.

Read with the document environment (python-docx/PyMuPDF); no clinical decision
is written, no production database is opened for mutation, no portal is used.
"""
import argparse
import csv
import hashlib
import json
import re
import statistics
import subprocess
from pathlib import Path
from zipfile import ZipFile

from docx import Document
import pymupdf

ROOT=Path(__file__).resolve().parents[1]
REV=ROOT/"revision/major_revision_20261006"
OUT=ROOT/"submission/bmc_bioinformatics_major_revision_20261019"
RECORD=REV/"final_review_20261019"


def sha(path):
    with path.open("rb") as f: return hashlib.file_digest(f,"sha256").hexdigest()


def comments(text):
    starts=list(re.finditer(r"^### (R[13]-[\w-]+)\.",text,re.M))
    result={}
    for i,m in enumerate(starts):
        body=text[m.end():starts[i+1].start() if i+1<len(starts) else len(text)]
        result[m[1]]=re.search(r"\*\*Comment\.\*\*\n\n(.*?)\n\n\*\*Response\.\*\*",body,re.S)[1]
    return result


def main():
    global OUT, RECORD
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir",type=Path,default=OUT)
    parser.add_argument("--record-dir",type=Path,default=RECORD)
    args=parser.parse_args()
    OUT,RECORD=args.output_dir,args.record_dir
    RECORD.mkdir(parents=True,exist_ok=True)
    baseline=subprocess.check_output(["rtk","proxy","git","show","d3d309a:revision/major_revision_20261006/rebuttal.md"],text=True,cwd=ROOT)
    response=(REV/"rebuttal.md").read_text()
    assert len(comments(response))==23 and comments(baseline)==comments(response)
    assert "> The figures were not in an accessible form. So I could not review this." in response
    assert "## Editor\n" in response
    assert len(list(csv.DictReader((REV/"reviewer_crosswalk.tsv").open(),delimiter="\t")))==25
    protected=json.loads((REV/"phase5_claude_review/baseline_verification.json").read_text())["sha256"]
    for name,h in protected.items(): assert sha(ROOT/name)==h,name
    old=ROOT/"submission/bmc_bioinformatics_major_revision_20261007"
    tagged_manifest=subprocess.check_output(["rtk","proxy","git","show","major-revision-author-review-20261007:submission/bmc_bioinformatics_major_revision_20261007/package_manifest.json"],cwd=ROOT)
    assert (old/"package_manifest.json").read_bytes()==tagged_manifest
    old_manifest=json.loads((old/"package_manifest.json").read_text())
    for name,m in old_manifest["artifacts"].items(): assert sha(old/name)==m["sha256"],name
    manifest=json.loads((OUT/"package_manifest.json").read_text())
    actual={str(p.relative_to(OUT)) for p in OUT.rglob("*") if p.is_file() and p.name not in ("package_manifest.json","package_verification.json")}
    assert actual==set(manifest["artifacts"]), actual ^ set(manifest["artifacts"])
    for name,h in manifest["exact_source_hashes"].items(): assert sha(ROOT/name)==h,name
    assert set(manifest["exact_source_hashes"])=={"manuscript/draft.md","manuscript/supplementary_methods.md","revision/major_revision_20261006/rebuttal.md","revision/major_revision_20261006/cover_letter_revision.md"}
    for name,m in manifest["artifacts"].items(): assert sha(OUT/name)==m["sha256"],name
    source_text="\n".join((ROOT/n).read_text() for n in manifest["exact_source_hashes"])
    for pattern in (r"await(?:s)? (?:case-by-case |individual )?human verification",r"has not yet been individually verified",r"have not yet.*verified by a human author",r"working draft",r"requires alignment with the final manuscript"):
        assert not re.search(pattern,source_text,re.I),pattern
    assert "Eligibility of the flagged cases remains under review" in response
    assert "None passed FDR < 0.05" in (ROOT/"manuscript/supplementary_methods.md").read_text()
    assert source_text.count("## S8.")==1
    body,reference_text=(ROOT/"manuscript/draft.md").read_text().split("## References\n",1)
    labels=[int(n) for n in re.findall(r"^\[(\d+)\]",reference_text,re.M)]
    assert labels==list(range(1,67))
    first=[]
    for m in re.finditer(r"\[(\d+(?:,\s*\d+)*)\]",body):
        values=[int(v) for v in re.split(r",\s*",m[1])]
        if not all(v in labels for v in values): continue
        for v in values:
            if v not in first: first.append(v)
    assert first==labels,"References must follow first main-text citation"
    ranks=list(csv.DictReader((REV/"temporal_validation_v2/results/case_rankings.tsv").open(),delimiter="\t"))
    temporal={}
    for s,top,ten,med in (("archived_five",2,5,46.23),("augmented_six",3,6,44.95)):
        cases=[r for r in ranks if r["scheme"]==s and r["clinical_status"]=="strong_novel_target"]
        coding=[r for r in cases if r["in_historical_universe"]=="True"]
        assert len(cases)==45 and len(coding)==41
        vals={k:sum(r[k]=="True" for r in coding) for k in ("top100","high","top1000","top10pct")}
        assert vals==dict(top100=0,high=0,top1000=top,top10pct=ten)
        assert round(statistics.median(float(r["percentile"]) for r in coding),2)==med
        temporal[s]=vals
    documents={}
    for name,drawings,tables in (("Manuscript",9,4),("Supplementary_Methods",1,2),("Response_to_Reviewers",0,0),("Cover_Letter",0,0)):
        doc=Document(OUT/f"{name}.docx")
        assert len(doc.inline_shapes)==drawings and len(doc.tables)==tables,(name,len(doc.inline_shapes),len(doc.tables))
        if drawings: assert all(s._inline.docPr.get("descr") for s in doc.inline_shapes)
        if name=="Manuscript":
            assert doc.sections[0]._sectPr.xpath(".//w:lnNumType")
            assert doc.styles["Normal"].paragraph_format.line_spacing==2.0
        pdf=pymupdf.open(OUT/f"{name}.pdf")
        text="\n".join(p.get_text() for p in pdf)
        assert "\ufffd" not in text and all(p.get_text().strip() for p in pdf)
        compact=re.sub(r"\s+","",text)
        for stale in ("Clinicalclassificationsawaitcase-by-casehumanverification", "hasnotyetbeenindividuallyverifiedbyahumanauthor", "workingdraft"):
            assert stale not in compact,(name,stale)
        if name=="Manuscript":
            assert manifest["current_author_review_release_tag"] in compact
            assert "Subsequenthumanreviewidentifiedunresolvedtarget-familyevidenceinthreeflaggedcases" in compact
            assert "BP/MF/CCcoveragewas41/42/44" in compact
        if name=="Supplementary_Methods":
            assert "S6.6Subsequenthumanauthorreview" in compact and "S8.ContemporaryGOclassification" in compact
            assert "passed362tests" in compact
        if name=="Response_to_Reviewers":
            assert "Thefigureswerenotinanaccessibleform" in compact
        images=sum(len(p.get_image_info()) for p in pdf)
        assert images>=drawings,(name,images)
        (RECORD/f"{name}_pdf_text.txt").write_text(text)
        documents[name]=dict(pages=len(pdf),embedded_images=images,word_tables=len(doc.tables),pdf_sha256=sha(OUT/f"{name}.pdf"))
    locations=list(csv.DictReader((OUT/"Page_location_index.tsv").open(),delimiter="\t"))
    assert {r["comment_id"] for r in locations}==set(comments(response))
    for r in locations:
        end=int(r["pdf_pages"].split("–")[-1])
        assert end<=documents[Path(r["document"]).stem]["pages"]
    for p in OUT.glob("Additional_file_*.zip"):
        assert p.stat().st_size<20_000_000,p
        with ZipFile(p) as z:
            assert z.testzip() is None
            for entry in json.loads(z.read("BUNDLE_MANIFEST.json")):
                payload=z.read(entry["path"])
                assert hashlib.sha256(payload).hexdigest()==entry["sha256"]
                assert sha(ROOT/entry["path"])==entry["sha256"],entry["path"]
            assert not any("principal_case_human_verification.tsv" in n or "Temporal_Case_Author_Review" in n or n.endswith(".inspect.ndjson") for n in z.namelist())
            if "Additional_file_13_" in p.name:
                assert "tests/test_revision_go.py" in z.namelist()
    figures=list((OUT/"figures").glob("*.pdf"))
    assert len(figures)==9
    assert all(p.stat().st_size<10_000_000 for p in (OUT/"figures").iterdir() if p.suffix in (".pdf",".png"))
    go=json.loads((REV/"go_classification_20261019/verification.json").read_text())
    assert go["status"]=="PASS" and sum(x["significant"] for x in go["independent_fisher_BH"])==0
    unchanged=subprocess.run(["rtk","proxy","git","diff","--exit-code","95662c7","--","revision/major_revision_20261006/temporal_validation_v2"],cwd=ROOT,capture_output=True,text=True)
    assert unchanged.returncode==0,"Frozen temporal tree changed from task-start checkpoint"
    report=dict(status="PASS",original_comments_preserved=23,response_topics=25,protected_baseline_hashes=len(protected),old_package_artifacts_unchanged=len(old_manifest["artifacts"]),current_source_hashes_exact=True,documents=documents,page_location_rows=len(locations),original_temporal_outcomes=temporal,GO_independent_tests=24597,additional_ZIP_size_limit_bytes=20_000_000,full_author_review_form_not_attached=True,limits=["Author approval of the revised version remains the purpose of this package","Flagged clinical eligibility remains under review; no new approved cohort asserted","Journal-generated PDF/reviewer file access can only be checked in the submission portal"])
    (RECORD/"verification.json").write_text(json.dumps(report,indent=2)+"\n")
    print(json.dumps(report,indent=2))


if __name__=="__main__": main()
